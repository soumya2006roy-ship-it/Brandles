from dataclasses import dataclass
from hashlib import blake2b
from io import BytesIO
import re
import unicodedata
import zipfile

import pymupdf
from docx import Document
from fastapi import HTTPException, UploadFile


MAX_SOURCE_BYTES = 20 * 1024 * 1024
CHUNK_CHARACTERS = 1400
CHUNK_OVERLAP = 180
EMBEDDING_DIMENSIONS = 1536


@dataclass(frozen=True)
class ExtractedBlock:
    text: str
    page_number: int | None = None
    location_label: str | None = None


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    return "".join(char for char in text if char in "\n\t" or unicodedata.category(char) not in {"Cc", "Cf"})


async def read_upload(upload: UploadFile) -> bytes:
    contents = await upload.read(MAX_SOURCE_BYTES + 1)
    if not contents:
        raise HTTPException(status_code=422, detail="The uploaded document is empty.")
    if len(contents) > MAX_SOURCE_BYTES:
        raise HTTPException(status_code=413, detail="Source documents must be 20 MB or smaller.")
    return contents


def identify_format(filename: str, contents: bytes) -> tuple[str, str]:
    extension = filename.rsplit(".", maxsplit=1)[-1].lower() if "." in filename else ""
    if contents.startswith(b"%PDF-"):
        if extension != "pdf":
            raise HTTPException(status_code=415, detail="The file contents are PDF but the filename extension is not .pdf.")
        return "pdf", "application/pdf"
    if contents.startswith(b"PK\x03\x04"):
        if extension != "docx" or not zipfile.is_zipfile(BytesIO(contents)):
            raise HTTPException(status_code=415, detail="Only valid DOCX documents are supported for ZIP-based files.")
        with zipfile.ZipFile(BytesIO(contents)) as archive:
            entries = archive.infolist()
            if "[Content_Types].xml" not in archive.namelist() or "word/document.xml" not in archive.namelist():
                raise HTTPException(status_code=415, detail="This ZIP file is not a valid DOCX document.")
            expanded_size = sum(entry.file_size for entry in entries)
            if expanded_size > 100 * 1024 * 1024 or expanded_size > max(len(contents), 1) * 100:
                raise HTTPException(status_code=413, detail="This DOCX document expands beyond the safe processing limit.")
        return "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if contents.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image_png", "image/png"
    if contents.startswith(b"\xff\xd8\xff"):
        return "image_jpeg", "image/jpeg"
    if len(contents) >= 12 and contents[:4] == b"RIFF" and contents[8:12] == b"WEBP":
        return "image_webp", "image/webp"
    if extension not in {"txt", "md"}:
        raise HTTPException(status_code=415, detail="Upload a PDF, DOCX, TXT, or Markdown file.")
    try:
        contents.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=415, detail="Text and Markdown documents must use UTF-8 encoding.") from exc
    return extension, "text/markdown" if extension == "md" else "text/plain"


def extract_document(file_type: str, contents: bytes) -> list[ExtractedBlock]:
    try:
        if file_type == "pdf":
            pdf = pymupdf.open(stream=contents, filetype="pdf")
            if pdf.is_encrypted:
                raise HTTPException(status_code=422, detail="Password-protected PDFs cannot be added as approved sources.")
            blocks = [ExtractedBlock(clean_text(page.get_text("text")), index + 1, f"Page {index + 1}") for index, page in enumerate(pdf)]
        elif file_type == "docx":
            document = Document(BytesIO(contents))
            lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            blocks = _split_blocks("\n".join(lines), "Paragraphs")
        else:
            blocks = _split_blocks(contents.decode("utf-8-sig"), "Lines")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The file could not be parsed. Check that it is a valid, readable document.") from exc
    normalized = [ExtractedBlock(clean_text(block.text).strip(), block.page_number, block.location_label) for block in blocks]
    normalized = [block for block in normalized if block.text]
    if not normalized:
        raise HTTPException(status_code=422, detail="No readable text was found in this document.")
    return normalized


def _split_blocks(text: str, prefix: str) -> list[ExtractedBlock]:
    lines = [line.strip() for line in clean_text(text).splitlines() if line.strip()]
    blocks: list[ExtractedBlock] = []
    current: list[str] = []
    length = 0
    start_line = 1
    for line_number, line in enumerate(lines, start=1):
        if length + len(line) > 5000 and current:
            blocks.append(ExtractedBlock("\n".join(current), None, f"{prefix} {start_line}–{line_number - 1}"))
            current, length, start_line = [], 0, line_number
        current.append(line)
        length += len(line) + 1
    if current:
        blocks.append(ExtractedBlock("\n".join(current), None, f"{prefix} {start_line}–{len(lines)}"))
    return blocks


def chunk_block(block: ExtractedBlock) -> list[tuple[str, int, int]]:
    text = re.sub(r"\s+", " ", block.text).strip()
    chunks: list[tuple[str, int, int]] = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_CHARACTERS, len(text))
        if end < len(text):
            boundary = text.rfind(" ", start + CHUNK_CHARACTERS // 2, end)
            if boundary > start:
                end = boundary
        chunks.append((text[start:end].strip(), start, end))
        if end == len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def local_hash_embedding(text: str) -> list[float]:
    """Deterministic lexical feature vector for local search, not a semantic model."""
    vector = [0.0] * EMBEDDING_DIMENSIONS
    tokens = re.findall(r"[\w%]+", text.casefold(), flags=re.UNICODE)
    features = tokens + [f"{left} {right}" for left, right in zip(tokens, tokens[1:])]
    for feature in features:
        digest = blake2b(feature.encode("utf-8"), digest_size=8).digest()
        value = int.from_bytes(digest, "big")
        index = value % EMBEDDING_DIMENSIONS
        vector[index] += 1.0 if value & (1 << 63) else -1.0
    norm = sum(value * value for value in vector) ** 0.5 or 1.0
    return [value / norm for value in vector]
